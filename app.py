import os
import io
import json
import re
import zipfile
import xml.etree.ElementTree as ET
from flask import Flask, render_template, request, send_file, jsonify
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

app = Flask(__name__)

PARTICIPANTS_FILE = os.path.join(os.path.dirname(__file__), "participants.json")
SETTINGS_FILE = os.path.join(os.path.dirname(__file__), "settings.json")

DEFAULT_SETTINGS = {
    "current_invoice_number": 1,
    "hourly_rate": 45.28,
    "provider": {
        "name": "Richa Patel",
        "business_name": "Richa Patel",
        "account_name": "Richa Patel",
        "abn": "88 475 952 165",
        "address": "8/41 McMinn Street, Darwin City NT 0800",
        "email_phone": "ripatel291202@gmail.com",
        "bank_name": "Westpac Bank",
        "bsb": "732-273",
        "account_number": "502413"
    },
    "customer": {
        "name": "Top End Support Collective",
        "attention": "Kerrie Toll",
        "abn": "",
        "address": "1 Palmerston cct",
        "reference": "Subcontract service claim"
    }
}

def load_settings():
    if not os.path.exists(SETTINGS_FILE):
        with open(SETTINGS_FILE, "w", encoding="utf-8") as f:
            json.dump(DEFAULT_SETTINGS, f, indent=2)
        return DEFAULT_SETTINGS
    try:
        with open(SETTINGS_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return DEFAULT_SETTINGS

def save_settings(data):
    with open(SETTINGS_FILE, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2)

DEFAULT_PARTICIPANTS = []

def load_participants():
    if not os.path.exists(PARTICIPANTS_FILE):
        with open(PARTICIPANTS_FILE, "w", encoding="utf-8") as f:
            json.dump(DEFAULT_PARTICIPANTS, f, indent=2)
        return DEFAULT_PARTICIPANTS
    try:
        with open(PARTICIPANTS_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return DEFAULT_PARTICIPANTS

def save_participants(data):
    with open(PARTICIPANTS_FILE, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2)

@app.route("/")
def index():
    participants = load_participants()
    settings = load_settings()
    return render_template("index.html", participants=participants, settings=settings)

@app.route("/api/settings", methods=["GET", "POST"])
def settings_api():
    if request.method == "POST":
        data = request.json
        save_settings(data)
        return jsonify({"status": "success", "settings": data})
    return jsonify(load_settings())

@app.route("/api/participants", methods=["GET", "POST"])
def participants_api():
    if request.method == "POST":
        data = request.json
        save_participants(data)
        return jsonify({"status": "success", "participants": data})
    return jsonify(load_participants())

@app.route("/api/participants/<pid>", methods=["DELETE"])
def delete_participant_api(pid):
    participants = load_participants()
    updated = [p for p in participants if p.get("id") != pid]
    save_participants(updated)
    return jsonify({"status": "success", "participants": updated})

# Shared Styling Constants
LIGHT_BLUE_HEADER = "D9E1F2"
NAVY_TABLE_HEADER = "104E8B"
LIGHT_GREEN_NOTE = "E2EFDA"
BORDER_COLOR = "D9D9D9"
DARK_TOTAL_BG = "1F4E79"
WHITE = "FFFFFF"

font_sec_hdr = Font(name="Aptos Narrow", size=11, bold=True, color="000000")
font_tbl_hdr = Font(name="Aptos Narrow", size=10, bold=True, color=WHITE)
font_lbl = Font(name="Aptos Narrow", size=10, bold=True, color="000000")
font_val = Font(name="Aptos Narrow", size=10, color="000000")
font_total_bold = Font(name="Aptos Narrow", size=11, bold=True, color=WHITE)
font_disclaimer = Font(name="Aptos Narrow", size=9, bold=True, color="333333")
font_links = Font(name="Aptos Narrow", size=8.5, italic=True, color="555555")

fill_sec_hdr = PatternFill(start_color=LIGHT_BLUE_HEADER, end_color=LIGHT_BLUE_HEADER, fill_type="solid")
fill_tbl_hdr = PatternFill(start_color=NAVY_TABLE_HEADER, end_color=NAVY_TABLE_HEADER, fill_type="solid")
fill_green_box = PatternFill(start_color=LIGHT_GREEN_NOTE, end_color=LIGHT_GREEN_NOTE, fill_type="solid")
fill_dark_total = PatternFill(start_color=DARK_TOTAL_BG, end_color=DARK_TOTAL_BG, fill_type="solid")

thin_border = Border(
    left=Side(style="thin", color=BORDER_COLOR),
    right=Side(style="thin", color=BORDER_COLOR),
    top=Side(style="thin", color=BORDER_COLOR),
    bottom=Side(style="thin", color=BORDER_COLOR)
)

align_left = Alignment(horizontal="left", vertical="center")
align_right = Alignment(horizontal="right", vertical="center")
align_center = Alignment(horizontal="center", vertical="center")
align_wrap = Alignment(horizontal="left", vertical="top", wrap_text=True)

def build_participant_sheet(ws, p_data, provider, customer, rates):
    ws.views.sheetView[0].showGridLines = True

    # Row 5: Section Headers
    ws.merge_cells("A5:C5")
    ws["A5"] = "FROM / SERVICE PROVIDER"
    ws["A5"].font = font_sec_hdr
    ws["A5"].fill = fill_sec_hdr
    ws["A5"].alignment = align_left

    ws.merge_cells("D5:F5")
    ws["D5"] = "BILL TO"
    ws["D5"].font = font_sec_hdr
    ws["D5"].fill = fill_sec_hdr
    ws["D5"].alignment = align_left

    pairs_row6_10 = [
        ("Provider", provider.get("name", "Richa Patel"), "Customer", customer.get("name", "Top End Support Collective")),
        ("Business / trading name", provider.get("business_name", "Richa Patel"), "Attention", customer.get("attention", "Kerrie Toll")),
        ("ABN", provider.get("abn", "88 475 952 165"), "Customer ABN", customer.get("abn", "")),
        ("Address", provider.get("address", "8/41 McMinn Street, Darwin City NT 0800"), "Address", customer.get("address", "1 Palmerston cct")),
        ("Email / phone", provider.get("email_phone", "ripatel291202@gmail.com"), "Reference", customer.get("reference", "Subcontract service claim")),
    ]

    for idx, (p_lbl, p_val, c_lbl, c_val) in enumerate(pairs_row6_10, start=6):
        ws[f"A{idx}"] = p_lbl
        ws[f"A{idx}"].font = font_lbl
        ws[f"B{idx}"] = p_val
        ws[f"B{idx}"].font = font_val
        ws.merge_cells(f"B{idx}:C{idx}")
        
        ws[f"D{idx}"] = c_lbl
        ws[f"D{idx}"].font = font_lbl
        ws[f"E{idx}"] = c_val
        ws[f"E{idx}"].font = font_val
        ws.merge_cells(f"E{idx}:F{idx}")

    inv_num = str(p_data.get("invoice_number", "1"))
    inv_date = p_data.get("invoice_date", "07 Sep 2026")
    inv_due = p_data.get("invoice_due", "21 Sep 2026")

    pairs_row12_16 = [
        ("Invoice number", inv_num, "Participant", p_data.get("name", "Participant")),
        ("Invoice date", inv_date, "NDIS number", str(p_data.get("ndis_number", ""))),
        ("Payment due", inv_due, "Participant address", str(p_data.get("address", ""))),
        ("Claim type", p_data.get("claim_type", "Standard"), "Support category", p_data.get("support_category", "Support Coordination")),
        ("GST treatment", p_data.get("gst_treatment", "GST-free NDIS support"), "Support item", str(p_data.get("support_item", "07_002_0106_8_3"))),
    ]

    for idx, (i_lbl, i_val, pt_lbl, pt_val) in enumerate(pairs_row12_16, start=12):
        ws[f"A{idx}"] = i_lbl
        ws[f"A{idx}"].font = font_lbl
        ws[f"B{idx}"] = i_val
        ws[f"B{idx}"].font = font_val
        ws.merge_cells(f"B{idx}:C{idx}")
        
        ws[f"D{idx}"] = pt_lbl
        ws[f"D{idx}"].font = font_lbl
        ws[f"E{idx}"] = pt_val
        ws[f"E{idx}"].font = font_val
        ws.merge_cells(f"E{idx}:F{idx}")

    for r in range(5, 17):
        if r == 11: continue
        for c in range(1, 7):
            ws.cell(row=r, column=c).border = thin_border

    table_headers = ["Service period", "NDIS support delivered", "Support item", "Minutes", "Hours", "Amount"]
    table_start_row = 18
    for c_i, h_text in enumerate(table_headers, start=1):
        cell = ws.cell(row=table_start_row, column=c_i, value=h_text)
        cell.font = font_tbl_hdr
        cell.fill = fill_tbl_hdr
        cell.alignment = align_center if c_i in [1, 3, 4] else (align_right if c_i in [5, 6] else align_left)
        cell.border = thin_border

    items = p_data.get("items", [])
    if not items:
        items = [{
            "service_period": "12 Aug 2026",
            "support_delivered": "Coordination of Supports - Level 2",
            "support_item": p_data.get("support_item", "07_002_0106_8_3"),
            "minutes": 60,
            "hours": 1.00,
            "amount": 45.28
        }]

    first_data_row = 19
    current_row = first_data_row
    cached_cells = {}
    subcontract_rate = float(rates.get("subcontract_rate", 45.28))
    tot_hrs = 0.0
    tot_amt = 0.0

    for item in items:
        ws.cell(row=current_row, column=1, value=item.get("service_period", "")).alignment = align_center
        ws.cell(row=current_row, column=2, value=item.get("support_delivered", "")).alignment = align_left
        ws.cell(row=current_row, column=3, value=item.get("support_item", "")).alignment = align_center
        
        mins = float(item.get("minutes", 0))
        hrs = round(mins / 60, 2)
        amt = round(hrs * subcontract_rate, 2)
        tot_hrs += hrs
        tot_amt += amt

        c_min = ws.cell(row=current_row, column=4, value=mins)
        c_min.alignment = align_center
        c_min.number_format = "0"
        
        c_hrs = ws.cell(row=current_row, column=5)
        c_hrs.value = f"=ROUND(D{current_row}/60, 2)"
        c_hrs.alignment = align_right
        c_hrs.number_format = "0.00"
        cached_cells[f"E{current_row}"] = hrs
        
        c_amt = ws.cell(row=current_row, column=6)
        c_amt.value = f"=ROUND(E{current_row}*$B$27, 2)"
        c_amt.alignment = align_right
        c_amt.number_format = "$#,##0.00"
        cached_cells[f"F{current_row}"] = amt

        for col_idx in range(1, 7):
            c_cell = ws.cell(row=current_row, column=col_idx)
            c_cell.font = font_val
            c_cell.border = thin_border

        current_row += 1

    last_data_row = current_row - 1
    subtotal_row = last_data_row + 2

    # Subtotal
    ws.merge_cells(f"D{subtotal_row}:E{subtotal_row}")
    ws[f"D{subtotal_row}"] = "Subtotal"
    ws[f"D{subtotal_row}"].font = font_lbl
    ws[f"D{subtotal_row}"].alignment = align_left
    
    subtot_c = ws[f"F{subtotal_row}"]
    subtot_c.value = f"=SUM(F{first_data_row}:F{last_data_row})"
    subtot_c.font = font_val
    subtot_c.alignment = align_right
    subtot_c.number_format = "$#,##0.00"
    cached_cells[f"F{subtotal_row}"] = round(tot_amt, 2)
    
    for c in range(4, 7):
        ws.cell(row=subtotal_row, column=c).border = thin_border

    # GST 0%
    gst_row = subtotal_row + 1
    ws.merge_cells(f"D{gst_row}:E{gst_row}")
    ws[f"D{gst_row}"] = "GST (0%)"
    ws[f"D{gst_row}"].font = font_lbl
    ws[f"D{gst_row}"].alignment = align_left
    
    gst_c = ws[f"F{gst_row}"]
    gst_c.value = 0.00
    gst_c.font = font_val
    gst_c.alignment = align_right
    gst_c.number_format = "$#,##0.00"
    for c in range(4, 7):
        ws.cell(row=gst_row, column=c).border = thin_border

    # Invoice total
    total_row = gst_row + 1
    ws.merge_cells(f"D{total_row}:E{total_row}")
    ws[f"D{total_row}"] = "Invoice total"
    ws[f"D{total_row}"].font = font_total_bold
    ws[f"D{total_row}"].fill = fill_dark_total
    ws[f"D{total_row}"].alignment = align_left
    
    tot_c = ws[f"F{total_row}"]
    tot_c.value = f"=F{subtotal_row}+F{gst_row}"
    tot_c.font = font_total_bold
    tot_c.fill = fill_dark_total
    tot_c.alignment = align_right
    tot_c.number_format = "$#,##0.00"
    cached_cells[f"F{total_row}"] = round(tot_amt, 2)
    for c in range(4, 7):
        ws.cell(row=total_row, column=c).border = thin_border

    # Rate and Payment Details
    rate_header_row = total_row + 3
    rate_cell_row = rate_header_row + 1

    for r_idx in range(first_data_row, last_data_row + 1):
        ws.cell(row=r_idx, column=6).value = f"=ROUND(E{r_idx}*$B${rate_cell_row}, 2)"

    ws.merge_cells(f"A{rate_header_row}:C{rate_header_row}")
    ws[f"A{rate_header_row}"] = "RATE AND PAYMENT DETAILS"
    ws[f"A{rate_header_row}"].font = font_sec_hdr
    ws[f"A{rate_header_row}"].fill = fill_sec_hdr
    ws[f"A{rate_header_row}"].alignment = align_left
    
    ws.merge_cells(f"D{rate_header_row}:F{rate_header_row}")
    ws[f"D{rate_header_row}"] = "PAYMENT NOTE"
    ws[f"D{rate_header_row}"].font = font_sec_hdr
    ws[f"D{rate_header_row}"].fill = fill_green_box
    ws[f"D{rate_header_row}"].alignment = align_left

    subcontract_rate = float(rates.get("subcontract_rate", 45.28))
    source_rate_desc = rates.get("source_rate_desc", "$45.28 / hr agreed rate")
    bank_bsb = provider.get("bsb", "732-273")
    bank_account = provider.get("account_number", "502413")
    bank_name = provider.get("bank_name", "Westpac Bank")
    account_name = provider.get("account_name", "Richa Patel")

    rate_rows = [
        ("Subcontract rate (per hour)", subcontract_rate, "$#,##0.00"),
        ("Source rate", source_rate_desc, "@"),
        ("Account name", account_name, "@"),
        ("BSB", str(bank_bsb), "@"),
        ("Account number", str(bank_account), "@")
    ]

    for offset, (lbl, val, n_fmt) in enumerate(rate_rows, start=1):
        r_i = rate_header_row + offset
        ws[f"A{r_i}"] = lbl
        ws[f"A{r_i}"].font = font_lbl
        ws[f"B{r_i}"] = val
        ws[f"B{r_i}"].font = font_val
        ws[f"B{r_i}"].number_format = n_fmt
        ws.merge_cells(f"B{r_i}:C{r_i}")
        for c in range(1, 4):
            ws[f"A{r_i}"].border = thin_border
            ws[f"B{r_i}"].border = thin_border
            ws[f"C{r_i}"].border = thin_border

    note_start_row = rate_header_row + 1
    note_end_row = rate_header_row + len(rate_rows)
    ws.merge_cells(f"D{note_start_row}:F{note_end_row}")
    
    remit_to = f"Please remit to {bank_name} using the account details shown." if bank_name else "Please remit using the account details shown."
    payment_note_text = (
        f"Payment reference: {inv_num}. Services are invoiced after delivery. {remit_to} "
        f"The hourly subcontract rate is ${subcontract_rate:.2f}/hr "
        f"and remains below the 2026-27 national NDIS price limit for Support Coordination Level 2."
    )
    ws[f"D{note_start_row}"] = payment_note_text
    ws[f"D{note_start_row}"].font = font_val
    ws[f"D{note_start_row}"].alignment = align_wrap
    for r in range(rate_header_row, note_end_row + 1):
        for c in range(4, 7):
            cell = ws.cell(row=r, column=c)
            cell.border = thin_border
            cell.fill = fill_green_box

    # Disclaimers
    disclaimer_row = note_end_row + 2
    ws.merge_cells(f"A{disclaimer_row}:F{disclaimer_row}")
    ws[f"A{disclaimer_row}"] = "IMPORTANT: Confirm the participant address and any applicable customer ABN before issue. Retain case notes or support logs evidencing the service dates and duration. This workbook is a billing document and does not replace service-delivery records."
    ws[f"A{disclaimer_row}"].font = font_disclaimer
    ws[f"A{disclaimer_row}"].alignment = align_wrap
    for c in range(1, 7):
        ws.cell(row=disclaimer_row, column=c).border = thin_border

    ref_row = disclaimer_row + 2
    ws.merge_cells(f"A{ref_row}:F{ref_row}")
    ws[f"A{ref_row}"] = "NDIS references: https://www.ndis.gov.au/providers/working-providers/reporting-and-recording-keeping/what-are-record-keeping-requirements | https://www.ndis.gov.au/providers/pricing-and-payments/pricing/pricing-arrangements"
    ws[f"A{ref_row}"].font = font_links
    ws[f"A{ref_row}"].alignment = align_left

    col_widths = {"A": 24, "B": 32, "C": 18, "D": 14, "E": 20, "F": 18}
    for col_l, w in col_widths.items():
        ws.column_dimensions[col_l].width = w

    return {
        "total_row": total_row,
        "subtotal_row": subtotal_row,
        "first_data_row": first_data_row,
        "last_data_row": last_data_row,
        "hours_col": "E",
        "amount_col": "F",
        "cached_cells": cached_cells,
        "total_hours": round(tot_hrs, 2),
        "subtotal": round(tot_amt, 2)
    }

def build_master_summary_sheet(ws_master, participants, provider, customer, rates, batch_info, p_sheet_results):
    ws_master.views.sheetView[0].showGridLines = True

    # Title Banner
    ws_master.merge_cells("A2:F2")
    ws_master["A2"] = "TAX INVOICE - CONSOLIDATED SUBCONTRACT CLAIM"
    ws_master["A2"].font = Font(name="Aptos Narrow", size=13, bold=True, color=NAVY_TABLE_HEADER)
    ws_master["A2"].alignment = align_left

    claim_period = batch_info.get("claim_period", "Fortnightly Subcontract Claim")
    ws_master.merge_cells("A3:F3")
    ws_master["A3"] = f"Master Billing Statement for Top End Support Collective - Period: {claim_period}"
    ws_master["A3"].font = Font(name="Aptos Narrow", size=10, italic=True, color="555555")
    ws_master["A3"].alignment = align_left

    # Row 5: Section Headers
    ws_master.merge_cells("A5:C5")
    ws_master["A5"] = "FROM / SERVICE PROVIDER"
    ws_master["A5"].font = font_sec_hdr
    ws_master["A5"].fill = fill_sec_hdr
    ws_master["A5"].alignment = align_left

    ws_master.merge_cells("D5:F5")
    ws_master["D5"] = "BILL TO / AGENCY OWNER"
    ws_master["D5"].font = font_sec_hdr
    ws_master["D5"].fill = fill_sec_hdr
    ws_master["D5"].alignment = align_left

    pairs_row6_10 = [
        ("Provider", provider.get("name", "Richa Patel"), "Customer", customer.get("name", "Top End Support Collective")),
        ("Business / trading name", provider.get("business_name", "Richa Patel"), "Attention", customer.get("attention", "Kerrie Toll")),
        ("ABN", provider.get("abn", "88 475 952 165"), "Customer ABN", customer.get("abn", "")),
        ("Address", provider.get("address", "8/41 McMinn Street, Darwin City NT 0800"), "Address", customer.get("address", "1 Palmerston cct")),
        ("Email / phone", provider.get("email_phone", "ripatel291202@gmail.com"), "Reference", batch_info.get("reference", "Fortnightly Subcontract Claim")),
    ]

    for idx, (p_lbl, p_val, c_lbl, c_val) in enumerate(pairs_row6_10, start=6):
        ws_master[f"A{idx}"] = p_lbl
        ws_master[f"A{idx}"].font = font_lbl
        ws_master[f"B{idx}"] = p_val
        ws_master[f"B{idx}"].font = font_val
        ws_master.merge_cells(f"B{idx}:C{idx}")

        ws_master[f"D{idx}"] = c_lbl
        ws_master[f"D{idx}"].font = font_lbl
        ws_master[f"E{idx}"] = c_val
        ws_master[f"E{idx}"].font = font_val
        ws_master.merge_cells(f"E{idx}:F{idx}")

        for c in range(1, 7):
            ws_master.cell(row=idx, column=c).border = thin_border

    # Row 12: Master Table Header
    master_headers = ["#", "Participant Name", "NDIS Number", "Invoice Ref", "Total Hours", "Total Claim Amount"]
    for c_i, h_text in enumerate(master_headers, start=1):
        c = ws_master.cell(row=12, column=c_i, value=h_text)
        c.font = font_tbl_hdr
        c.fill = fill_tbl_hdr
        c.alignment = align_center if c_i in [1, 3, 4] else (align_right if c_i in [5, 6] else align_left)
        c.border = thin_border

    start_summary_row = 13
    current_summary_row = start_summary_row
    master_cached_cells = {}

    for p_idx, res_item in enumerate(p_sheet_results, start=1):
        p = res_item["participant"]
        sheet_title = res_item["title"]
        res = res_item["res"]

        p_hrs = float(res.get("total_hours", 0.0))
        p_amt = float(res.get("subtotal", 0.0))

        inv_ref = f"Inv #{p.get('invoice_number', p_idx)}"
        
        ws_master.cell(row=current_summary_row, column=1, value=p_idx).alignment = align_center
        ws_master.cell(row=current_summary_row, column=2, value=p.get("name", f"Participant {p_idx}")).alignment = align_left
        ws_master.cell(row=current_summary_row, column=3, value=str(p.get("ndis_number", ""))).alignment = align_center
        ws_master.cell(row=current_summary_row, column=4, value=inv_ref).alignment = align_center
        
        # Link hours and amount dynamically to participant worksheet
        c_hrs = ws_master.cell(row=current_summary_row, column=5)
        c_hrs.value = f"=ROUND(SUM('{sheet_title}'!E{res['first_data_row']}:E{res['last_data_row']}), 2)"
        c_hrs.alignment = align_right
        c_hrs.number_format = "0.00"
        master_cached_cells[f"E{current_summary_row}"] = round(p_hrs, 2)

        c_amt = ws_master.cell(row=current_summary_row, column=6)
        c_amt.value = f"='{sheet_title}'!F{res['total_row']}"
        c_amt.alignment = align_right
        c_amt.number_format = "$#,##0.00"
        master_cached_cells[f"F{current_summary_row}"] = round(p_amt, 2)

        for c in range(1, 7):
            cell = ws_master.cell(row=current_summary_row, column=c)
            cell.font = font_val
            cell.border = thin_border

        current_summary_row += 1

    last_summary_row = max(start_summary_row, current_summary_row - 1)
    subtotal_row = last_summary_row + 2

    grand_total_hours = sum(float(r["res"].get("total_hours", 0.0)) for r in p_sheet_results)
    grand_total_amount = sum(float(r["res"].get("subtotal", 0.0)) for r in p_sheet_results)

    # Totals block in Master Sheet
    ws_master.merge_cells(f"C{subtotal_row}:D{subtotal_row}")
    ws_master[f"C{subtotal_row}"] = "Total Hours & Subtotal"
    ws_master[f"C{subtotal_row}"].font = font_lbl
    ws_master[f"C{subtotal_row}"].alignment = align_left
    
    tot_hrs_c = ws_master[f"E{subtotal_row}"]
    tot_hrs_c.value = f"=SUM(E{start_summary_row}:E{last_summary_row})"
    tot_hrs_c.font = font_lbl
    tot_hrs_c.alignment = align_right
    tot_hrs_c.number_format = "0.00"
    master_cached_cells[f"E{subtotal_row}"] = round(grand_total_hours, 2)

    subtot_c = ws_master[f"F{subtotal_row}"]
    subtot_c.value = f"=SUM(F{start_summary_row}:F{last_summary_row})"
    subtot_c.font = font_lbl
    subtot_c.alignment = align_right
    subtot_c.number_format = "$#,##0.00"
    master_cached_cells[f"F{subtotal_row}"] = round(grand_total_amount, 2)

    for c in range(3, 7):
        ws_master.cell(row=subtotal_row, column=c).border = thin_border

    # GST 0%
    gst_row = subtotal_row + 1
    ws_master.merge_cells(f"C{gst_row}:E{gst_row}")
    ws_master[f"C{gst_row}"] = "GST (0%)"
    ws_master[f"C{gst_row}"].font = font_lbl
    
    gst_c = ws_master[f"F{gst_row}"]
    gst_c.value = 0.00
    gst_c.font = font_val
    gst_c.alignment = align_right
    gst_c.number_format = "$#,##0.00"
    for c in range(3, 7):
        ws_master.cell(row=gst_row, column=c).border = thin_border

    # Final Bill Total Due
    final_total_row = gst_row + 1
    ws_master.merge_cells(f"C{final_total_row}:E{final_total_row}")
    ws_master[f"C{final_total_row}"] = "FINAL BILL TOTAL DUE"
    ws_master[f"C{final_total_row}"].font = font_total_bold
    ws_master[f"C{final_total_row}"].fill = fill_dark_total
    ws_master[f"C{final_total_row}"].alignment = align_left

    final_tot_c = ws_master[f"F{final_total_row}"]
    final_tot_c.value = f"=F{subtotal_row}+F{gst_row}"
    final_tot_c.font = font_total_bold
    final_tot_c.fill = fill_dark_total
    final_tot_c.alignment = align_right
    final_tot_c.number_format = "$#,##0.00"
    master_cached_cells[f"F{final_total_row}"] = round(grand_total_amount, 2)
    for c in range(3, 7):
        ws_master.cell(row=final_total_row, column=c).border = thin_border

    # Remittance Box in Master Sheet
    remit_header_row = final_total_row + 3
    ws_master.merge_cells(f"A{remit_header_row}:C{remit_header_row}")
    ws_master[f"A{remit_header_row}"] = "REMITTANCE & PAYMENT INSTRUCTIONS"
    ws_master[f"A{remit_header_row}"].font = font_sec_hdr
    ws_master[f"A{remit_header_row}"].fill = fill_sec_hdr

    ws_master.merge_cells(f"D{remit_header_row}:F{remit_header_row}")
    ws_master[f"D{remit_header_row}"] = "OWNER SUMMARY NOTE"
    ws_master[f"D{remit_header_row}"].font = font_sec_hdr
    ws_master[f"D{remit_header_row}"].fill = fill_green_box

    subcontract_rate = float(rates.get("subcontract_rate", 45.28))
    bank_bsb = provider.get("bsb", "732-273")
    bank_account = provider.get("account_number", "502413")
    bank_name = provider.get("bank_name", "Westpac Bank")
    account_name = provider.get("account_name", "Richa Patel")

    pay_rows = [
        ("Agreed Subcontract Rate", subcontract_rate, "$#,##0.00"),
        ("Account Name", account_name, "@"),
        ("Bank Name", bank_name, "@"),
        ("BSB Number", str(bank_bsb), "@"),
        ("Account Number", str(bank_account), "@")
    ]
    for offset, (lbl, val, n_fmt) in enumerate(pay_rows, start=1):
        r_i = remit_header_row + offset
        ws_master[f"A{r_i}"] = lbl
        ws_master[f"A{r_i}"].font = font_lbl
        ws_master[f"B{r_i}"] = val
        ws_master[f"B{r_i}"].font = font_val
        ws_master[f"B{r_i}"].number_format = n_fmt
        ws_master.merge_cells(f"B{r_i}:C{r_i}")
        for c in range(1, 4):
            ws_master.cell(row=r_i, column=c).border = thin_border

    note_start_row = remit_header_row + 1
    note_end_row = remit_header_row + len(pay_rows)
    ws_master.merge_cells(f"D{note_start_row}:F{note_end_row}")
    master_note = (
        f"Consolidated Subcontract Claim for Top End Support Collective ({len(participants)} participants). "
        f"Please make a single direct transfer of the Final Bill Total to {bank_name} using the account details shown. "
        f"Individual participant claim breakdowns are attached in the following workbook tabs for your NDIS PRODA / PACE claiming records."
    )
    ws_master[f"D{note_start_row}"] = master_note
    ws_master[f"D{note_start_row}"].font = font_val
    ws_master[f"D{note_start_row}"].alignment = align_wrap
    for r in range(remit_header_row, note_end_row + 1):
        for c in range(4, 7):
            cell = ws_master.cell(row=r, column=c)
            cell.border = thin_border
            cell.fill = fill_green_box

    # Disclaimers
    disclaimer_row = note_end_row + 2
    ws_master.merge_cells(f"A{disclaimer_row}:F{disclaimer_row}")
    ws_master[f"A{disclaimer_row}"] = "IMPORTANT: Retain case notes or support logs evidencing service dates and duration for each participant listed. This statement is a consolidated billing document and does not replace service-delivery records."
    ws_master[f"A{disclaimer_row}"].font = font_disclaimer
    ws_master[f"A{disclaimer_row}"].alignment = align_wrap
    for c in range(1, 7):
        ws_master.cell(row=disclaimer_row, column=c).border = thin_border

    master_col_widths = {"A": 6, "B": 30, "C": 18, "D": 18, "E": 16, "F": 20}
    for col_l, w in master_col_widths.items():
        ws_master.column_dimensions[col_l].width = w

    return {
        "master_cached_cells": master_cached_cells
    }

def inject_cached_values_into_xlsx(xlsx_bytes, cached_values_map):
    zf_in = zipfile.ZipFile(io.BytesIO(xlsx_bytes), 'r')
    
    wb_tree = ET.fromstring(zf_in.read('xl/workbook.xml'))
    rels_tree = ET.fromstring(zf_in.read('xl/_rels/workbook.xml.rels'))
    
    sheet_to_rid = {}
    for elem in wb_tree.iter():
        if elem.tag.endswith('sheet'):
            name = elem.attrib.get('name')
            for k, v in elem.attrib.items():
                if k.endswith('id'):
                    sheet_to_rid[name] = v
                    
    rid_to_target = {}
    for elem in rels_tree.iter():
        if elem.tag.endswith('Relationship'):
            rid = elem.attrib.get('Id')
            target = elem.attrib.get('Target', '').lstrip('/')
            if not target.startswith('xl/'):
                target = 'xl/' + target
            rid_to_target[rid] = target
            
    sheet_to_file = {s: rid_to_target[r] for s, r in sheet_to_rid.items() if r in rid_to_target}
    file_to_sheet = {f: s for s, f in sheet_to_file.items()}
    
    out_bytes = io.BytesIO()
    zf_out = zipfile.ZipFile(out_bytes, 'w', compression=zipfile.ZIP_DEFLATED)
    
    for item in zf_in.infolist():
        data = zf_in.read(item.filename)
        if item.filename in file_to_sheet:
            s_name = file_to_sheet[item.filename]
            if s_name in cached_values_map:
                cell_vals = cached_values_map[s_name]
                root = ET.fromstring(data)
                for c in root.iter():
                    if c.tag.endswith('}c'):
                        coord = c.attrib.get('r')
                        if coord in cell_vals:
                            val = cell_vals[coord]
                            if isinstance(val, (int, float)):
                                c.attrib.pop('t', None)
                            v_elem = None
                            for child in c:
                                if child.tag.endswith('}v'):
                                    v_elem = child
                                    break
                            if v_elem is None:
                                ns = c.tag.split('}')[0] + '}' if '}' in c.tag else ''
                                v_elem = ET.SubElement(c, f'{ns}v')
                            v_elem.text = str(val)
                data = ET.tostring(root, encoding='utf-8', xml_declaration=True)
        zf_out.writestr(item, data)
        
    zf_out.close()
    zf_in.close()
    out_bytes.seek(0)
    return out_bytes.getvalue()

@app.route("/api/export-excel", methods=["POST"])
def export_excel():
    data = request.json or {}
    provider = data.get("provider", {})
    customer = data.get("customer", {})
    rates = data.get("rates", {})
    batch_info = {
        "claim_period": data.get("claim_period", "Fortnightly Subcontract Claim"),
        "reference": data.get("reference", customer.get("reference", "Subcontract service claim"))
    }
    
    participants = data.get("participants", [])
    is_batch = (data.get("mode") != "single") and (len(participants) > 1 or data.get("is_batch") or data.get("mode") == "batch")

    wb = openpyxl.Workbook()
    wb.calculation.fullCalcOnLoad = True

    all_cached_values = {}

    if is_batch and participants:
        # Multi-tab workbook: Tab 1 = Master Bill, Tabs 2..N = Each participant
        ws_master = wb.active
        ws_master.title = "Master Bill - Owner Summary"

        p_sheet_results = []
        used_sheet_names = set(["Master Bill - Owner Summary"])

        for p_idx, p in enumerate(participants, start=1):
            raw_title = p.get("name", f"Participant_{p_idx}").strip()
            clean_title = re.sub(r'[\\/*?:\[\]]', '_', raw_title)[:28]
            title = clean_title
            c = 2
            while title in used_sheet_names:
                title = f"{clean_title[:24]} ({c})"
                c += 1
            used_sheet_names.add(title)

            ws_p = wb.create_sheet(title=title)
            res = build_participant_sheet(ws_p, p, provider, customer, rates)
            p_sheet_results.append({
                "title": title,
                "participant": p,
                "res": res
            })
            all_cached_values[title] = res.get("cached_cells", {})

        m_res = build_master_summary_sheet(ws_master, participants, provider, customer, rates, batch_info, p_sheet_results)
        all_cached_values["Master Bill - Owner Summary"] = m_res.get("master_cached_cells", {})
        
        # Safe filename
        safe_period = re.sub(r'[^a-zA-Z0-9_-]', '_', batch_info["claim_period"]).strip('_')
        download_filename = f"Final_Bill_Top_End_Support_Collective_{safe_period}.xlsx"

    else:
        # Single Invoice Mode
        ws = wb.active
        p_data = data.get("participant") or (participants[0] if participants else {})
        p_name = p_data.get("name", "Invoice")
        ws.title = re.sub(r'[\\/*?:\[\]]', '_', p_name)[:30]
        
        res = build_participant_sheet(ws, p_data, provider, customer, rates)
        all_cached_values[ws.title] = res.get("cached_cells", {})
        
        inv_num = str(p_data.get("invoice_number", "1"))
        clean_pname = "".join(c for c in p_name if c.isalnum() or c in " _-").strip()
        download_filename = f"Invoice_{inv_num}_{clean_pname}.xlsx"

    output = io.BytesIO()
    wb.save(output)
    raw_bytes = output.getvalue()

    final_bytes = inject_cached_values_into_xlsx(raw_bytes, all_cached_values)

    return send_file(
        io.BytesIO(final_bytes),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=download_filename
    )

if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=port, debug=False)
